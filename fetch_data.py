#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
数据拉取脚本：从 SQL Server 和抖音小店拉取销售数据并合并
"""

import os
import pandas as pd
import pyodbc
from datetime import datetime
from openpyxl.styles import Font, Alignment, numbers
from config import get_db_config


def fetch_pos_data(db_config, target_date, status=80):
    """
    从 SQL Server 拉取门店零售数据

    Args:
        db_config: 数据库配置字典
        target_date: 目标日期 (YYYY-MM-DD)
        status: 单据状态（默认 80）

    Returns:
        DataFrame: POS 销售数据
    """
    sql = f"""
    SELECT
        v.EntityCode,
        v.DocumentDate,
        d.ItemCode,
        d.Qty,
        v.TotalReceivableAmt,
        v.TotalDeductionAmt
    FROM {db_config['database']}.dbo.RT_POSVoucher v
    INNER JOIN {db_config['database']}.dbo.RT_POSVoucherDetail d
        ON v.POSVoucherID = d.POSVoucherID
    WHERE v.DocumentDate = '{target_date}'
        AND v.Status = {status}
    """
    
    print(f"正在从 SQL Server 拉取数据 (日期: {target_date}, 状态: {status})...")
    print("连接数据库...")
    
    # 使用 pyodbc 连接（pymssql 与 SQL Server 2014 存在兼容性问题）
    conn_str = f"DRIVER={{SQL Server}};SERVER={db_config['server']},{db_config['port']};DATABASE={db_config['database']};UID={db_config['user']};PWD={db_config['password']}"
    conn = pyodbc.connect(conn_str, timeout=120)
    
    print("数据库连接成功，正在执行查询...")
    
    try:
        df = pd.read_sql(sql, conn)
        print(f"SQL Server 数据拉取完成，共 {len(df)} 条记录")
        
        # 添加渠道标识
        df['Channel'] = 'POS'
        
        return df
    finally:
        conn.close()

def fetch_douyin_data(file_path):
    """
    从抖音小店 Excel 文件读取销售数据
    
    Args:
        file_path: Excel 文件路径
    
    Returns:
        DataFrame: 抖音销售数据
    """
    print(f"正在读取抖音小店数据: {file_path}")
    
    if not os.path.exists(file_path):
        print(f"警告: 文件不存在 {file_path}")
        return pd.DataFrame()
    
    df = pd.read_excel(file_path)
    print(f"抖音数据读取完成，共 {len(df)} 条记录")
    
    # 添加渠道标识
    df['Channel'] = '抖音'
    
    return df


def unify_format(pos_df, douyin_df):
    """
    统一两个渠道的数据格式
    
    Args:
        pos_df: POS 数据 DataFrame
        douyin_df: 抖音数据 DataFrame
    
    Returns:
        DataFrame: 合并后的数据
    """
    # 统一 POS 数据格式
    if not pos_df.empty:
        pos_df = pos_df.rename(columns={
            'EntityCode': 'StoreCode',
            'DocumentDate': 'SaleDate',
            'ItemCode': 'ProductCode',
            'Qty': 'SaleQty',
            'TotalReceivableAmt': 'ReceivableAmt',
            'TotalDeductionAmt': 'DeductionAmt'
        })
        pos_df['SaleDate'] = pd.to_datetime(pos_df['SaleDate']).dt.strftime('%Y-%m-%d')
        # 门店编码转为字符串，避免 Excel 中显示为科学计数法
        pos_df['StoreCode'] = pos_df['StoreCode'].astype(str)
    
    # 统一抖音数据格式（根据实际 Excel 列名映射）
    if not douyin_df.empty:
        column_mapping = {
            '门店编号': 'StoreCode',
            '销售日期': 'SaleDate',
            '款号': 'ProductCode',
            '数量': 'SaleQty',
            '金额': 'ReceivableAmt',
            '成本': 'DeductionAmt'
        }
        douyin_df = douyin_df.rename(columns=column_mapping)

        if 'SaleDate' in douyin_df.columns:
            douyin_df['SaleDate'] = pd.to_datetime(douyin_df['SaleDate']).dt.strftime('%Y-%m-%d')
        
        # 门店编码转为字符串
        if 'StoreCode' in douyin_df.columns:
            douyin_df['StoreCode'] = douyin_df['StoreCode'].astype(str)
    
    # 确保两个渠道的列顺序一致
    final_columns = ['StoreCode', 'SaleDate', 'ProductCode', 'SaleQty', 'ReceivableAmt', 'DeductionAmt', 'Channel']
    
    for df in [pos_df, douyin_df]:
        if not df.empty:
            for col in final_columns:
                if col not in df.columns:
                    df[col] = None
    
    if not pos_df.empty:
        pos_df = pos_df[final_columns]
    if not douyin_df.empty:
        douyin_df = douyin_df[final_columns]
    
    # 合并数据
    merged_df = pd.concat([pos_df, douyin_df], ignore_index=True)
    print(f"数据合并完成，共 {len(merged_df)} 条记录")
    
    return merged_df


def save_to_excel(df, output_dir='data'):
    """
    保存数据到 Excel 文件（带格式优化）

    Args:
        df: 数据 DataFrame
        output_dir: 输出目录
    """
    os.makedirs(output_dir, exist_ok=True)

    today = datetime.now().strftime('%Y_%m_%d')
    filename = f"{output_dir}/raw_{today}.xlsx"

    # 使用 openpyxl 直接写入（避免 pandas 类型推断）
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '销售数据'

    # 列名中文映射
    header_names = {
        'StoreCode': '门店编码',
        'SaleDate': '销售日期',
        'ProductCode': '商品编码',
        'SaleQty': '销售数量',
        'ReceivableAmt': '应收金额',
        'DeductionAmt': '扣减金额',
        'Channel': '渠道'
    }

    # 写表头
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_names.get(col_name, col_name))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 写数据行
    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        for col_idx, col_name in enumerate(df.columns, 1):
            value = row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # 门店编码列设为文本
            if col_name == 'StoreCode' and value is not None:
                cell.value = str(value)
                cell.number_format = numbers.FORMAT_TEXT

            # 日期列格式
            elif col_name == 'SaleDate' and value:
                cell.number_format = 'YYYY-MM-DD'

            # 金额列格式
            elif col_name in ('ReceivableAmt', 'DeductionAmt') and value is not None:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')

            # 数量列居中
            elif col_name == 'SaleQty' and value is not None:
                cell.alignment = Alignment(horizontal='center')

            # 渠道列居中
            elif col_name == 'Channel':
                cell.alignment = Alignment(horizontal='center')

    # 设置列宽
    for col_idx, col_name in enumerate(df.columns, 1):
        max_len = len(header_names.get(col_name, col_name))
        for row_idx in range(2, len(df) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                content_len = sum(2 if ord(c) > 127 else 1 for c in str(cell_value))
                max_len = max(max_len, content_len)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    # 冻结首行
    ws.freeze_panes = 'A2'

    # 添加筛选器
    ws.auto_filter.ref = f'A1:{get_column_letter(len(df.columns))}{len(df) + 1}'

    wb.save(filename)

    print(f"数据已保存至: {filename}")

    return filename


def save_to_csv(df, output_dir='data'):
    """
    保存数据到 CSV 文件（备用）

    Args:
        df: 数据 DataFrame
        output_dir: 输出目录
    """
    os.makedirs(output_dir, exist_ok=True)

    today = datetime.now().strftime('%Y_%m_%d')
    filename = f"{output_dir}/raw_{today}.csv"

    df.to_csv(filename, index=False, encoding='utf-8-sig')
    print(f"CSV 数据已保存至: {filename}")

    return filename


def main():
    """主函数"""
    print("=" * 60)
    print("开始执行数据拉取任务")
    print("=" * 60)
    
    # 加载数据库配置
    db_config = get_db_config()
    
    # 定义查询时间（2026年7月31日，状态80）
    target_date = '2026-07-31'
    
    # 拉取 POS 数据
    pos_df = fetch_pos_data(db_config, target_date)
    
    # 拉取抖音数据
    douyin_file = 'data/douyin.xlsx'
    douyin_df = fetch_douyin_data(douyin_file)
    
    # 统一格式并合并
    merged_df = unify_format(pos_df, douyin_df)
    
    # 保存到 Excel
    output_file = save_to_excel(merged_df, 'data')
    
    print("=" * 60)
    print(f"任务完成！输出文件: {output_file}")
    print("=" * 60)


if __name__ == '__main__':
    main()
